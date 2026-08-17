from __future__ import annotations

import csv
import os
import tempfile
from datetime import date
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .result_schema import RESULT_FIELDS, serialize_csv_value, validate_result
from .presentation import EXCEL_FIELDS, format_excel_text


EXPORT_FIELDS = ("condition", "source", "model", "reasoning_effort", *RESULT_FIELDS)
CONDITION_ORDER = {"WITH_SKILL": 0, "WITHOUT_SKILL": 1}


def build_with_skill_workbook(result: dict[str, Any]) -> bytes:
    """Build a one-row user-facing XLSX while retaining typed amount and date cells."""
    validate_result(result, str(result.get("thread_id", "")))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "A 결과"
    sheet.append(EXCEL_FIELDS)
    row: list[Any] = []
    for field in EXCEL_FIELDS:
        if field == "날짜":
            row.append(date.fromisoformat(result[field]))
        elif field == "금액":
            row.append(result[field])
        else:
            row.append(format_excel_text(field, result))
    sheet.append(row)
    sheet.freeze_panes = "A2"
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    widths = (28, 13, 34, 16, 10, 48, 56, 56, 56)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width
    for cell in sheet[2]:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet["B2"].number_format = "yyyy-mm-dd"
    sheet["D2"].number_format = "#,##0"
    sheet.auto_filter.ref = f"A1:I2"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class ABResultExcelRepository:
    """Upsert validated, deterministically ordered Excel-compatible CSV rows."""

    def __init__(self, output_path: Path, allowed_root: Path) -> None:
        self.output_path = output_path.resolve()
        self.allowed_root = allowed_root.resolve()
        if self.output_path.suffix.lower() != ".csv":
            raise ValueError("A/B export path must use the .csv extension.")
        if self.allowed_root not in self.output_path.parents:
            raise ValueError("A/B export must stay inside the configured experiment root.")

    def append_result(
        self,
        condition: str,
        source: str,
        thread_id: str,
        model: str,
        reasoning_effort: str,
        result: dict[str, Any],
    ) -> Path:
        return self.upsert_results([{
            "condition": condition,
            "source": source,
            "thread_id": thread_id,
            "model": model,
            "reasoning_effort": reasoning_effort,
            "result": result,
        }])

    def upsert_results(self, records: list[dict[str, Any]]) -> Path:
        existing = self._read_existing()
        return self._save_records(records, existing)

    def replace_results(self, records: list[dict[str, Any]]) -> Path:
        """Replace the result set after validating every supplied record."""
        return self._save_records(records, {})

    def _save_records(
        self,
        records: list[dict[str, Any]],
        existing: dict[tuple[str, str], dict[str, Any]],
    ) -> Path:
        for record in records:
            condition = str(record["condition"])
            if condition not in CONDITION_ORDER:
                raise ValueError(f"unsupported A/B condition: {condition}")
            result = validate_result(record["result"], str(record["thread_id"]))
            row = {
                "condition": condition,
                "source": str(record["source"]),
                "model": str(record["model"]),
                "reasoning_effort": str(record["reasoning_effort"]),
                **{field: serialize_csv_value(field, result[field]) for field in RESULT_FIELDS},
            }
            existing[(result["thread_id"], condition)] = row
        rows = sorted(
            existing.values(),
            key=lambda row: (row["thread_id"], CONDITION_ORDER[row["condition"]]),
        )
        self._write_atomic(rows)
        return self.output_path

    def _read_existing(self) -> dict[tuple[str, str], dict[str, Any]]:
        if not self.output_path.exists():
            return {}
        with self.output_path.open(encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            if tuple(reader.fieldnames or ()) != EXPORT_FIELDS:
                raise ValueError("existing CSV schema does not match the normalized export schema")
            rows = list(reader)
        return {(row["thread_id"], row["condition"]): row for row in rows}

    def _write_atomic(self, rows: list[dict[str, Any]]) -> None:
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        descriptor, temp_name = tempfile.mkstemp(
            prefix=f".{self.output_path.name}.", suffix=".tmp", dir=self.output_path.parent,
        )
        try:
            with os.fdopen(descriptor, "w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=EXPORT_FIELDS)
                writer.writeheader()
                writer.writerows(rows)
            Path(temp_name).replace(self.output_path)
        except Exception:
            Path(temp_name).unlink(missing_ok=True)
            raise
