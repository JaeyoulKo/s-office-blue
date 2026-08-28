from __future__ import annotations

import json
import os
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


READY_STATUS = "READY_FOR_APPROVAL_REVIEW"
WORKSHEET_TITLE = "승인 가능 이메일"
EXCEL_FIELDS = (
    "저장 일시", "case_id", "message_id", "thread_id", "발신자", "수신자",
    "수신 일시", "제목", "본문", "첨부파일", "문서 유형", "검토 상태",
    "검토 사유", "승인 안내", "승인 URL",
)
EXCEL_WIDTHS = (20, 28, 28, 28, 28, 34, 22, 46, 72, 36, 24, 30, 60, 60, 46)


class PurchaseApprovalExcelStore:
    """승인 가능한 구매 이메일만 지정된 runtime XLSX에 누적한다."""

    def __init__(self, output_path: Path, allowed_root: Path) -> None:
        self.output_path = output_path.resolve()
        self.allowed_root = allowed_root.resolve()
        if self.output_path.suffix.lower() != ".xlsx":
            raise ValueError("누적 Excel 파일은 .xlsx 확장자를 사용해야 합니다.")
        if self.allowed_root != self.output_path.parent and self.allowed_root not in self.output_path.parents:
            raise ValueError("누적 Excel 파일은 지정된 runtime 경로 안에 있어야 합니다.")

    def append_approved_email(
        self,
        email: dict[str, Any],
        review: dict[str, Any],
        *,
        saved_at: datetime | None = None,
    ) -> tuple[int, bool]:
        if review.get("review_status") != READY_STATUS:
            raise ValueError("승인 가능으로 판정된 이메일만 Excel에 저장할 수 있습니다.")
        record_key = _record_key(email)
        workbook, sheet = self._load_or_create()
        existing_row = _find_record_row(sheet, record_key)
        if existing_row is not None:
            workbook.close()
            return existing_row, False
        row_number = _append_row(
            sheet, email, review, saved_at=saved_at or datetime.now(timezone.utc)
        )
        self._save_atomic(workbook)
        return row_number, True

    def read_bytes(self) -> bytes | None:
        if not self.output_path.is_file():
            return None
        workbook, _ = self._load_existing()
        workbook.close()
        return self.output_path.read_bytes()

    def _load_or_create(self) -> tuple[Workbook, Any]:
        return self._load_existing() if self.output_path.exists() else _new_workbook()

    def _load_existing(self) -> tuple[Workbook, Any]:
        try:
            workbook = load_workbook(self.output_path)
        except Exception as exc:
            raise ValueError(f"기존 누적 Excel 파일을 열 수 없습니다: {exc}") from exc
        if WORKSHEET_TITLE not in workbook.sheetnames:
            workbook.close()
            raise ValueError(f"기존 Excel에 '{WORKSHEET_TITLE}' sheet가 없습니다.")
        sheet = workbook[WORKSHEET_TITLE]
        header = tuple(
            sheet.cell(row=1, column=index).value
            for index in range(1, len(EXCEL_FIELDS) + 1)
        )
        if header != EXCEL_FIELDS:
            workbook.close()
            raise ValueError("기존 Excel 헤더가 현재 승인 가능 이메일 schema와 다릅니다.")
        return workbook, sheet

    def _save_atomic(self, workbook: Workbook) -> None:
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        descriptor, temp_name = tempfile.mkstemp(
            prefix=f".{self.output_path.stem}.", suffix=".xlsx", dir=self.output_path.parent
        )
        os.close(descriptor)
        temp_path = Path(temp_name)
        try:
            workbook.save(temp_path)
            workbook.close()
            temp_path.replace(self.output_path)
        except Exception:
            workbook.close()
            temp_path.unlink(missing_ok=True)
            raise


def _new_workbook() -> tuple[Workbook, Any]:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = WORKSHEET_TITLE
    sheet.append(EXCEL_FIELDS)
    sheet.freeze_panes = "A2"
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for index, width in enumerate(EXCEL_WIDTHS, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width
    sheet.auto_filter.ref = "A1:O1"
    return workbook, sheet


def _append_row(
    sheet: Any,
    email: dict[str, Any],
    review: dict[str, Any],
    *,
    saved_at: datetime,
) -> int:
    guidance = review.get("approval_guidance")
    guidance_summary = guidance.get("summary") if isinstance(guidance, dict) else guidance
    guidance_url = guidance.get("url") if isinstance(guidance, dict) else None
    attachments = email.get("attachments")
    attachment_names: list[str] = []
    if isinstance(attachments, list):
        for item in attachments:
            if isinstance(item, dict):
                attachment_names.append(str(item.get("file_name") or item.get("name") or "").strip())
            elif item:
                attachment_names.append(str(item).strip())
    row = (
        saved_at.replace(tzinfo=None),
        _safe_excel_text(email.get("case_id")),
        _safe_excel_text(email.get("message_id")),
        _safe_excel_text(email.get("thread_id")),
        _safe_excel_text(email.get("sender") or email.get("from")),
        _safe_excel_text(_join_values(email.get("recipients") or email.get("to"))),
        _safe_excel_text(email.get("received_at")),
        _safe_excel_text(email.get("subject")),
        _safe_excel_text(email.get("body")),
        _safe_excel_text("\n".join(name for name in attachment_names if name)),
        _safe_excel_text(review.get("document_type")),
        READY_STATUS,
        _safe_excel_text(review.get("status_reason")),
        _safe_excel_text(guidance_summary),
        _safe_excel_text(guidance_url or email.get("approval_url")),
    )
    sheet.append(row)
    row_number = sheet.max_row
    for cell in sheet[row_number]:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.cell(row=row_number, column=1).number_format = "yyyy-mm-dd hh:mm:ss"
    sheet.auto_filter.ref = f"A1:O{row_number}"
    return row_number


def _record_key(email: dict[str, Any]) -> str:
    for field in ("message_id", "case_id"):
        value = str(email.get(field) or "").strip()
        if value:
            return f"{field}:{value}"
    thread_id = str(email.get("thread_id") or "").strip()
    subject = str(email.get("subject") or "").strip()
    if thread_id or subject:
        return f"thread_subject:{thread_id}|{subject}"
    raise ValueError("중복 판별에 필요한 이메일 식별자가 없습니다.")


def _find_record_row(sheet: Any, record_key: str) -> int | None:
    for row_number in range(2, sheet.max_row + 1):
        existing = {
            "case_id": sheet.cell(row=row_number, column=2).value,
            "message_id": sheet.cell(row=row_number, column=3).value,
            "thread_id": sheet.cell(row=row_number, column=4).value,
            "subject": sheet.cell(row=row_number, column=8).value,
        }
        try:
            if _record_key(existing) == record_key:
                return row_number
        except ValueError:
            continue
    return None


def _join_values(value: Any) -> str:
    if isinstance(value, (list, tuple)):
        return ", ".join(str(item) for item in value if item not in (None, ""))
    return str(value or "")


def _safe_excel_text(value: Any) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, (dict, list, tuple)):
        text = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    else:
        text = str(value)
    return f"'{text}" if text.startswith(("=", "+", "-", "@")) else text
