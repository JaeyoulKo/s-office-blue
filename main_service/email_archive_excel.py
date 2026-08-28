from __future__ import annotations

import hashlib
import json
import os
import tempfile
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .archive_result import ArchiveResultFormatError, validate_archive_result
from .skill_registry import PROJECT_ROOT


MASTER_EXCEL_PATH = (
    PROJECT_ROOT
    / "experiments"
    / "instances"
    / "email_archive_exports"
    / "email_archive_results.xlsx"
)
WORKSHEET_TITLE = "A 결과"
METADATA_TITLE = "_office_blue_meta"
EXCEL_FIELDS = (
    "발신자",
    "날짜",
    "Topic",
    "금액",
    "통화",
    "Business Impact",
    "Thread 진행 중 변경된 내용",
    "결정된 내용",
    "Open Item",
)
EMPTY_VALUE = "확인된 내용 없음"
EXCEL_WIDTHS = (28, 13, 34, 16, 10, 48, 56, 56, 56)
HIDDEN_FIELDS = frozenset(
    {"thread_id", "message_id", "model", "reasoning_effort", "skill", "적용 규칙"}
)


def archive_save_candidates(
    records: dict[str, dict[str, Any]], order: list[str]
) -> tuple[list[tuple[dict[str, Any], dict[str, Any]]], int]:
    """현재 결과 영역의 성공한 아카이빙 결과만 Excel 저장 후보로 고른다."""
    successful: list[tuple[dict[str, Any], dict[str, Any]]] = []
    failed = 0
    for case_id in order:
        record = records.get(case_id)
        if record is None:
            continue  # 선택·분석되지 않은 inbox 항목은 저장 후보가 아니다.
        if record.get("status") != "ok":
            failed += 1
            continue
        result = record.get("result")
        parsed = getattr(result, "parsed", None)
        email = record.get("email")
        if not isinstance(email, dict):
            failed += 1
            continue
        try:
            canonical = validate_archive_result(parsed)
        except ArchiveResultFormatError:
            failed += 1
            continue
        successful.append((email, canonical))
    return successful, failed


def natural_language_items(value: Any, *, estimated: bool = False) -> list[str]:
    """중첩 결과를 화면과 Excel에 쓸 중복 없는 자연어 문장 목록으로 바꾼다."""
    items: list[str] = []
    seen: set[str] = set()

    def add(text: Any, prefix: str = "") -> None:
        if text in (None, ""):
            return
        if isinstance(text, bool):
            candidate = "예" if text else "아니요"
        elif isinstance(text, (int, float)):
            candidate = f"{text:,}"
        else:
            candidate = str(text).strip()
        if not candidate:
            return
        semantic_text = candidate.removeprefix("확인 필요: ").removeprefix("예상: ")
        if semantic_text in seen:
            return
        seen.add(semantic_text)
        items.append(f"{prefix}{semantic_text}")

    if isinstance(value, dict):
        for key, nested in value.items():
            if key in HIDDEN_FIELDS:
                continue
            if key == "confirmed":
                for item in natural_language_items(nested):
                    add(item)
            elif key == "estimated":
                for item in natural_language_items(nested):
                    add(item, "확인 필요: ")
            elif key == "items":
                for item in natural_language_items(nested):
                    add(item)
            elif key in {"field", "from", "to"}:
                continue
            else:
                for item in natural_language_items(nested, estimated=estimated):
                    add(item, "확인 필요: " if estimated else "")
    elif isinstance(value, (list, tuple, set)):
        for nested in value:
            if isinstance(nested, dict) and any(key in nested for key in ("from", "to")):
                label = str(nested.get("field") or "내용").strip()
                if label in HIDDEN_FIELDS:
                    continue
                before = _scalar_text(nested.get("from"))
                after = _scalar_text(nested.get("to"))
                if before != EMPTY_VALUE and after != EMPTY_VALUE:
                    add(f"{label}이(가) {before}에서 {after}(으)로 변경되었습니다.")
                elif after != EMPTY_VALUE:
                    add(f"{label}이(가) {after}(으)로 확정되었습니다.")
                elif before != EMPTY_VALUE:
                    add(f"{label}의 이전 값은 {before}입니다.")
            else:
                for item in natural_language_items(nested, estimated=estimated):
                    add(item, "확인 필요: " if estimated else "")
    else:
        add(value, "확인 필요: " if estimated else "")
    return items


def display_text(value: Any) -> str:
    items = natural_language_items(value)
    return "\n".join(items) if items else EMPTY_VALUE


def excel_text(value: Any) -> str:
    items = natural_language_items(value)
    return "\n".join(items) if items else EMPTY_VALUE


def formatted_amount(amount: Any, currency: Any = None) -> str:
    number = _numeric_amount(amount)
    if number is None:
        return EMPTY_VALUE
    text = f"{number:,}"
    return f"{text} {currency}" if isinstance(currency, str) and currency else text


def result_version_key(email: dict[str, Any], result: dict[str, Any]) -> str:
    thread_id = str(email.get("thread_id") or result.get("thread_id") or "")
    latest_date = str(email.get("received_at") or result.get("날짜") or "")
    normalized = {
        field: result.get(field)
        for field in (*EXCEL_FIELDS, "thread_id")
    }
    raw = json.dumps(
        {"thread_id": thread_id, "latest_date": latest_date, "result": normalized},
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        default=str,
    ).encode("utf-8")
    return hashlib.sha256(raw).hexdigest()


@dataclass(frozen=True)
class SaveResult:
    added: bool
    row_count: int


@dataclass(frozen=True)
class BatchSaveResult:
    added_count: int
    duplicate_count: int
    row_count: int


class EmailArchiveExcelStore:
    def __init__(self, output_path: Path = MASTER_EXCEL_PATH) -> None:
        self.output_path = output_path
        if self.output_path.suffix.lower() != ".xlsx":
            raise ValueError("Master Excel은 .xlsx 확장자를 사용해야 합니다.")

    def append(self, email: dict[str, Any], result: dict[str, Any]) -> SaveResult:
        saved = self.append_many([(email, result)])
        return SaveResult(added=saved.added_count == 1, row_count=saved.row_count)

    def append_many(
        self, records: list[tuple[dict[str, Any], dict[str, Any]]]
    ) -> BatchSaveResult:
        """여러 성공 결과를 한 번의 원자적 workbook 교체로 누적한다."""
        for _email, result in records:
            validate_archive_result(result)
        workbook, sheet, metadata = self._load_or_create()
        saved_keys = {
            str(metadata.cell(row=row, column=1).value or "")
            for row in range(2, metadata.max_row + 1)
        }
        added = duplicates = 0
        for email, result in records:
            version_key = result_version_key(email, result)
            if version_key in saved_keys:
                duplicates += 1
                continue
            sheet.append(_excel_row(result))
            row_number = sheet.max_row
            for cell in sheet[row_number]:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            sheet.cell(row=row_number, column=2).number_format = "yyyy-mm-dd"
            sheet.cell(row=row_number, column=4).number_format = "#,##0"
            metadata.append((version_key, row_number))
            saved_keys.add(version_key)
            added += 1
        sheet.auto_filter.ref = f"A1:I{sheet.max_row}"
        row_count = max(sheet.max_row - 1, 0)
        if added:
            self._save_atomic(workbook)
        else:
            workbook.close()
        return BatchSaveResult(added, duplicates, row_count)

    def read_bytes(self) -> bytes | None:
        if not self.output_path.is_file():
            return None
        return self.output_path.read_bytes()

    def _load_or_create(self) -> tuple[Workbook, Any, Any]:
        if self.output_path.is_file():
            workbook = load_workbook(self.output_path)
            if WORKSHEET_TITLE not in workbook.sheetnames:
                workbook.close()
                raise ValueError(f"기존 Excel에 '{WORKSHEET_TITLE}' worksheet가 없습니다.")
            sheet = workbook[WORKSHEET_TITLE]
            header = tuple(
                sheet.cell(row=1, column=index).value
                for index in range(1, len(EXCEL_FIELDS) + 1)
            )
            if header != EXCEL_FIELDS:
                workbook.close()
                raise ValueError("기존 Master Excel의 열 형식이 다릅니다.")
        else:
            workbook, sheet = _new_workbook()
        metadata = _metadata_sheet(workbook)
        return workbook, sheet, metadata

    def _save_atomic(self, workbook: Workbook) -> None:
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        descriptor, temp_name = tempfile.mkstemp(
            prefix=f".{self.output_path.stem}.",
            suffix=".xlsx",
            dir=self.output_path.parent,
        )
        os.close(descriptor)
        temp_path = Path(temp_name)
        try:
            workbook.save(temp_path)
            workbook.close()
            temp_path.replace(self.output_path)
        except Exception:
            workbook.close()
            temp_path.unlink(missing_ok=True)
            raise


def _new_workbook() -> tuple[Workbook, Any]:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = WORKSHEET_TITLE
    sheet.append(EXCEL_FIELDS)
    sheet.freeze_panes = "A2"
    fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for index, width in enumerate(EXCEL_WIDTHS, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width
    sheet.auto_filter.ref = "A1:I1"
    return workbook, sheet


def _metadata_sheet(workbook: Workbook) -> Any:
    if METADATA_TITLE in workbook.sheetnames:
        sheet = workbook[METADATA_TITLE]
    else:
        sheet = workbook.create_sheet(METADATA_TITLE)
        sheet.append(("version_key", "archive_row"))
    sheet.sheet_state = "hidden"
    return sheet


def _excel_row(result: dict[str, Any]) -> list[Any]:
    return [
        _safe_excel_text(display_text(result.get("발신자"))),
        _excel_date(result.get("날짜")),
        _safe_excel_text(display_text(result.get("Topic"))),
        _numeric_amount(result.get("금액")),
        _safe_excel_text(display_text(result.get("통화"))),
        _safe_excel_text(excel_text(result.get("Business Impact"))),
        _safe_excel_text(excel_text(result.get("Thread 진행 중 변경된 내용"))),
        _safe_excel_text(excel_text(result.get("결정된 내용"))),
        _safe_excel_text(excel_text(result.get("Open Item"))),
    ]


def _excel_date(value: Any) -> date | str:
    text = str(value or "").strip()
    try:
        return date.fromisoformat(text[:10])
    except ValueError:
        try:
            return datetime.fromisoformat(text.replace("Z", "+00:00")).date()
        except ValueError:
            return EMPTY_VALUE


def _numeric_amount(value: Any) -> int | float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return value
    text = str(value).replace(",", "").strip()
    try:
        number = float(text)
    except ValueError:
        return None
    return int(number) if number.is_integer() else number


def _scalar_text(value: Any) -> str:
    if value in (None, ""):
        return EMPTY_VALUE
    if isinstance(value, bool):
        return "예" if value else "아니요"
    if isinstance(value, (int, float)):
        return f"{value:,}"
    return str(value).strip() or EMPTY_VALUE


def _safe_excel_text(value: str) -> str:
    return f"'{value}" if value.startswith(("=", "+", "-", "@")) else value
