from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest


APP = Path(__file__).parents[1] / "streamlit_app.py"


def _app(tmp_path: Path, monkeypatch) -> tuple[AppTest, Path]:
    workbook = tmp_path / "archive.xlsx"
    monkeypatch.setenv("ARCHIVE_REPOSITORY_ROOT", str(tmp_path))
    monkeypatch.setenv("ARCHIVE_WORKBOOK", "archive.xlsx")
    app = AppTest.from_file(str(APP), default_timeout=30).run()
    app = app.radio[0].set_value("Development fixtures").run()
    return app, workbook


def _button(app: AppTest, label: str):
    return next(button for button in app.button if button.label == label)


def test_live_gmail_is_default_and_requires_inbox_selection(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setenv("ARCHIVE_REPOSITORY_ROOT", str(tmp_path))
    monkeypatch.setenv("ARCHIVE_WORKBOOK", "archive.xlsx")
    app = AppTest.from_file(str(APP), default_timeout=30).run()

    assert app.radio[0].value == "Live Gmail"
    assert _button(app, "Gmail 새로고침")
    assert _button(app, "Analyze Archive").disabled is True


def test_preview_does_not_write_and_save_creates_workbook(tmp_path: Path, monkeypatch) -> None:
    app, workbook = _app(tmp_path, monkeypatch)

    app = _button(app, "Analyze Archive").click().run()
    assert not app.exception
    assert not workbook.exists()
    assert any("new_record_ready" in item.value for item in app.success)

    app = _button(app, "Save Archive").click().run()
    assert not app.exception
    assert workbook.exists()
    assert any("Archive 저장 완료" in item.value for item in app.success)
    assert _button(app, "Save Archive").disabled is True

    book = load_workbook(workbook, read_only=True, data_only=True)
    try:
        assert book.sheetnames == [
            "Records",
            "Action Items",
            "Attachments",
            "Source Emails",
            "Change History",
        ]
    finally:
        book.close()


def test_update_and_ambiguous_match_ui(tmp_path: Path, monkeypatch) -> None:
    app, workbook = _app(tmp_path, monkeypatch)
    app = _button(app, "Analyze Archive").click().run()
    app = _button(app, "Save Archive").click().run()
    original = workbook.read_bytes()

    app = app.selectbox[0].set_value("Project Blue · thread update").run()
    app = _button(app, "Analyze Archive").click().run()
    assert any("record_update_ready" in item.value for item in app.info)
    assert any(subheader.value == "Change History" for subheader in app.subheader)

    app = app.selectbox[0].set_value("Project Blue · ambiguous candidate").run()
    app = _button(app, "Analyze Archive").click().run()
    assert any("additional confirmation required" in item.value for item in app.warning)
    assert _button(app, "Save Archive").disabled is True
    assert any("Candidate" in expander.label for expander in app.expander)
    assert workbook.read_bytes() == original
