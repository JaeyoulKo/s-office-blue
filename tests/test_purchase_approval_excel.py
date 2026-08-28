from __future__ import annotations

import importlib.util
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
MODULE_PATH = (
    PROJECT_ROOT
    / "experiments"
    / "ablation"
    / "purchase-email-review"
    / "approval_excel.py"
)
SPEC = importlib.util.spec_from_file_location("purchase_approval_excel", MODULE_PATH)
assert SPEC is not None and SPEC.loader is not None
approval_excel = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(approval_excel)

EXCEL_FIELDS = approval_excel.EXCEL_FIELDS
PurchaseApprovalExcelStore = approval_excel.PurchaseApprovalExcelStore
READY_STATUS = approval_excel.READY_STATUS
WORKSHEET_TITLE = approval_excel.WORKSHEET_TITLE


def email(**extra):
    return {
        "case_id": "case-1",
        "message_id": "message-1",
        "thread_id": "thread-1",
        "sender": "requester@example.invalid",
        "recipients": ["buyer@example.invalid"],
        "received_at": "2026-08-22T10:30:00+09:00",
        "subject": "구매 승인 요청",
        "body": "구매 목적과 비용 근거가 모두 포함된 본문",
        "attachments": [{"file_name": "quote.xlsx"}],
        "approval_url": "https://approval.example.invalid/PR-1",
        **extra,
    }


def review(**extra):
    return {
        "document_type": "ARIBA_SPEND_REQUEST",
        "review_status": READY_STATUS,
        "status_reason": "필수 네 항목이 모두 충족되었습니다.",
        "approval_guidance": {
            "summary": "승인 페이지에서 최종 확인하세요.",
            "url": "https://approval.example.invalid/PR-1",
        },
        **extra,
    }


class PurchaseApprovalExcelStoreTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.path = self.root / "exports" / "approved_purchase_emails.xlsx"
        self.store = PurchaseApprovalExcelStore(self.path, self.root)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_appends_ready_email_with_fixed_schema_and_types(self):
        row_number, created = self.store.append_approved_email(
            email(), review(), saved_at=datetime(2026, 8, 22, 1, 30, 0)
        )
        self.assertEqual((row_number, created), (2, True))
        workbook = load_workbook(self.path)
        sheet = workbook[WORKSHEET_TITLE]
        self.assertEqual(tuple(cell.value for cell in sheet[1]), EXCEL_FIELDS)
        self.assertEqual(sheet.cell(2, 3).value, "message-1")
        self.assertEqual(sheet.cell(2, 9).value, "구매 목적과 비용 근거가 모두 포함된 본문")
        self.assertEqual(sheet.cell(2, 12).value, READY_STATUS)
        self.assertIsInstance(sheet.cell(2, 1).value, datetime)
        workbook.close()

    def test_rejects_non_ready_result_without_creating_file(self):
        with self.assertRaisesRegex(ValueError, "승인 가능"):
            self.store.append_approved_email(
                email(), review(review_status="REQUEST_CLARIFICATION")
            )
        self.assertFalse(self.path.exists())

    def test_same_message_is_not_appended_twice(self):
        self.assertEqual(self.store.append_approved_email(email(), review()), (2, True))
        self.assertEqual(self.store.append_approved_email(email(), review()), (2, False))
        workbook = load_workbook(self.path)
        self.assertEqual(workbook[WORKSHEET_TITLE].max_row, 2)
        workbook.close()

    def test_formula_like_email_text_is_escaped(self):
        self.store.append_approved_email(email(subject='=HYPERLINK("bad")'), review())
        workbook = load_workbook(self.path, data_only=False)
        self.assertEqual(workbook[WORKSHEET_TITLE].cell(2, 8).value, "'=HYPERLINK(\"bad\")")
        workbook.close()

    def test_rejects_output_outside_allowed_root(self):
        with self.assertRaisesRegex(ValueError, "runtime 경로"):
            PurchaseApprovalExcelStore(self.root.parent / "outside.xlsx", self.root)

    def test_rejects_existing_workbook_with_wrong_header(self):
        self.path.parent.mkdir(parents=True)
        workbook = Workbook()
        workbook.active.title = WORKSHEET_TITLE
        workbook.active.append(["잘못된 헤더"])
        workbook.save(self.path)
        workbook.close()
        with self.assertRaisesRegex(ValueError, "헤더"):
            self.store.append_approved_email(email(), review())
