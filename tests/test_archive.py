from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from archive.analyzer import FixtureArchiveAnalyzer
from archive.cli import main as cli_main
from archive.gmail_adapter import GmailUnavailableError, UnavailableGmailGateway
from archive.models import EmailThread, ThreadAnalysis
from archive.service import ArchiveService
from archive.storage import ExcelRecordRepository

FIXTURES = Path(__file__).parent / "fixtures"


def load_thread(name: str) -> EmailThread:
    return EmailThread.model_validate_json((FIXTURES / name).read_text(encoding="utf-8"))


def load_analysis(name: str) -> ThreadAnalysis:
    return ThreadAnalysis.model_validate_json((FIXTURES / name).read_text(encoding="utf-8"))


def service(workbook: Path, analysis: ThreadAnalysis) -> ArchiveService:
    return ArchiveService(
        FixtureArchiveAnalyzer(analysis),
        ExcelRecordRepository(workbook, workbook.parent),
    )


def test_thread_is_sorted_chronologically() -> None:
    thread = load_thread("thread_update.json")
    assert [message.message_id for message in thread.messages] == ["msg-001", "msg-002"]


def test_missing_values_use_unknown_or_empty_lists() -> None:
    thread = EmailThread.model_validate(
        {
            "thread_id": "t",
            "messages": [
                {
                    "message_id": "m",
                    "thread_id": "t",
                    "sent_at": "2026-08-07T00:00:00Z",
                }
            ],
        }
    )
    assert thread.messages[0].subject == "unknown"
    assert thread.messages[0].recipients == []
    assert thread.messages[0].attachments == []


def test_empty_thread_is_rejected() -> None:
    with pytest.raises(ValidationError):
        EmailThread.model_validate({"thread_id": "t", "messages": []})


def test_attachment_metadata_is_preserved() -> None:
    attachment = load_thread("thread_initial.json").messages[0].attachments[0]
    assert attachment.file_name == "quote.pdf"
    assert attachment.access_status == "unverified"
    assert attachment.summary == "unknown"


def test_new_record_is_created_and_saved_to_all_sheets(tmp_path: Path) -> None:
    workbook = tmp_path / "archive.xlsx"
    result = service(workbook, load_analysis("analysis_initial.json")).process(
        load_thread("thread_initial.json"), write=True
    )
    assert result.result_status == "saved"
    assert result.storage.operation == "create"
    assert workbook.exists()
    book = load_workbook(workbook, data_only=True)
    try:
        assert book.sheetnames == [
            "Records",
            "Action Items",
            "Attachments",
            "Source Emails",
            "Change History",
        ]
        assert book["Records"].max_row == 2
        assert book["Action Items"].max_row == 2
        assert book["Attachments"].max_row == 2
        assert book["Source Emails"].max_row == 2
    finally:
        book.close()


def test_same_thread_does_not_create_duplicate_record(tmp_path: Path) -> None:
    workbook = tmp_path / "archive.xlsx"
    current_service = service(workbook, load_analysis("analysis_initial.json"))
    first = current_service.process(load_thread("thread_initial.json"), write=True)
    second = current_service.process(load_thread("thread_initial.json"), write=True)
    records = ExcelRecordRepository(workbook, tmp_path).list_records()
    assert second.storage.operation == "update"
    assert len(records) == 1
    assert records[0].record_id == first.storage.saved_record_id


def test_new_unrelated_thread_is_new_record(tmp_path: Path) -> None:
    workbook = tmp_path / "archive.xlsx"
    initial = load_thread("thread_initial.json")
    service(workbook, load_analysis("analysis_initial.json")).process(initial, write=True)
    other = initial.model_copy(deep=True)
    other.thread_id = "thread-unrelated"
    other.messages[0].thread_id = "thread-unrelated"
    other.messages[0].message_id = "msg-other"
    other.messages[0].subject = "Completely unrelated topic"
    other.messages[0].reference_ids = ["REF-OTHER"]
    analysis = load_analysis("analysis_initial.json").model_copy(deep=True)
    analysis.subject = "Completely unrelated topic"
    analysis.requester = "other@example.com"
    analysis.participants = ["other@example.com"]
    result = service(workbook, analysis).process(other, write=False)
    assert result.record_match.decision == "new_record"


def test_existing_record_update_and_due_date_change_history(tmp_path: Path) -> None:
    workbook = tmp_path / "archive.xlsx"
    initial_result = service(workbook, load_analysis("analysis_initial.json")).process(
        load_thread("thread_initial.json"), write=True
    )
    update_result = service(workbook, load_analysis("analysis_update.json")).process(
        load_thread("thread_update.json"), write=True
    )
    records = ExcelRecordRepository(workbook, tmp_path).list_records()
    assert update_result.result_status == "saved"
    assert update_result.storage.operation == "update"
    assert update_result.storage.saved_record_id == initial_result.storage.saved_record_id
    assert len(records) == 1
    assert records[0].due_dates == ["2026-08-15"]

    book = load_workbook(workbook, data_only=True)
    try:
        history = list(book["Change History"].iter_rows(min_row=2, values_only=True))
        due_changes = [row for row in history if str(row[2]).endswith(".due_date")]
        assert len(due_changes) == 1
        assert due_changes[0][4] == "2026-08-10"
        assert due_changes[0][5] == "2026-08-15"
        assert due_changes[0][7] == "msg-002"
    finally:
        book.close()


def test_change_history_is_appended_not_deleted(tmp_path: Path) -> None:
    workbook = tmp_path / "archive.xlsx"
    service(workbook, load_analysis("analysis_initial.json")).process(
        load_thread("thread_initial.json"), write=True
    )
    service(workbook, load_analysis("analysis_update.json")).process(
        load_thread("thread_update.json"), write=True
    )
    before = _sheet_data_rows(workbook, "Change History")

    completed_analysis = load_analysis("analysis_update.json").model_copy(deep=True)
    completed_analysis.action_items[0].status = "completed"
    completed_analysis.current_status = "completed"
    completed_thread = load_thread("thread_update.json").model_copy(deep=True)
    completed_thread.messages.append(
        completed_thread.messages[-1].model_copy(
            update={
                "message_id": "msg-003",
                "sent_at": datetime.fromisoformat("2026-08-09T10:00:00+09:00"),
                "body_text": "검토 완료",
            }
        )
    )
    completed_thread = EmailThread.model_validate(completed_thread.model_dump(mode="json"))
    service(workbook, completed_analysis).process(completed_thread, write=True)
    after = _sheet_data_rows(workbook, "Change History")
    assert after > before


def test_ambiguous_match_blocks_write_and_returns_candidates(tmp_path: Path) -> None:
    workbook = tmp_path / "archive.xlsx"
    service(workbook, load_analysis("analysis_initial.json")).process(
        load_thread("thread_initial.json"), write=True
    )
    before = workbook.read_bytes()

    ambiguous_thread = load_thread("thread_initial.json").model_copy(deep=True)
    ambiguous_thread.thread_id = "thread-ambiguous"
    ambiguous_thread.messages[0].thread_id = "thread-ambiguous"
    ambiguous_thread.messages[0].message_id = "msg-ambiguous"
    ambiguous_thread.messages[0].reference_ids = []
    result = service(workbook, load_analysis("analysis_initial.json")).process(
        ambiguous_thread, write=True
    )
    assert result.result_status == "additional_confirmation_required"
    assert result.record_match.candidates
    assert result.storage.executed is False
    assert workbook.read_bytes() == before


def test_without_write_workbook_is_not_created_or_modified(tmp_path: Path) -> None:
    workbook = tmp_path / "archive.xlsx"
    result = service(workbook, load_analysis("analysis_initial.json")).process(
        load_thread("thread_initial.json"), write=False
    )
    assert result.result_status == "new_record_ready"
    assert result.storage.executed is False
    assert not workbook.exists()


def test_without_write_existing_workbook_is_unchanged(tmp_path: Path) -> None:
    workbook = tmp_path / "archive.xlsx"
    service(workbook, load_analysis("analysis_initial.json")).process(
        load_thread("thread_initial.json"), write=True
    )
    before = workbook.read_bytes()
    result = service(workbook, load_analysis("analysis_update.json")).process(
        load_thread("thread_update.json"), write=False
    )
    assert result.result_status == "record_update_ready"
    assert workbook.read_bytes() == before


def test_repository_rejects_workbook_outside_configured_root(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="inside the repository"):
        ExcelRecordRepository(tmp_path.parent / "outside.xlsx", tmp_path)


def test_write_failure_does_not_damage_existing_workbook(tmp_path: Path) -> None:
    workbook = tmp_path / "archive.xlsx"
    service(workbook, load_analysis("analysis_initial.json")).process(
        load_thread("thread_initial.json"), write=True
    )
    before = workbook.read_bytes()

    class FailingValidationRepository(ExcelRecordRepository):
        def _validate_saved_file(self, path: Path, record_id: str) -> None:
            raise RuntimeError("simulated validation failure")

    failing_service = ArchiveService(
        FixtureArchiveAnalyzer(load_analysis("analysis_update.json")),
        FailingValidationRepository(workbook, tmp_path),
    )
    with pytest.raises(RuntimeError, match="simulated validation failure"):
        failing_service.process(load_thread("thread_update.json"), write=True)
    assert workbook.read_bytes() == before
    assert not list(tmp_path.glob(".archive.*.xlsx"))


def test_cli_preview_does_not_write(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr("archive.cli.REPOSITORY_ROOT", tmp_path)
    workbook = tmp_path / "archive.xlsx"
    exit_code = cli_main(
        [
            "--input",
            str(FIXTURES / "thread_initial.json"),
            "--analysis-fixture",
            str(FIXTURES / "analysis_initial.json"),
            "--workbook",
            str(workbook),
        ]
    )
    assert exit_code == 0
    assert not workbook.exists()
    assert '"result_status": "new_record_ready"' in capsys.readouterr().out


def test_gmail_gateway_reports_unavailable() -> None:
    with pytest.raises(GmailUnavailableError, match="No Gmail MCP"):
        UnavailableGmailGateway().fetch_thread("thread-id")


def _sheet_data_rows(workbook: Path, sheet_name: str) -> int:
    book = load_workbook(workbook, data_only=True)
    try:
        return max(0, book[sheet_name].max_row - 1)
    finally:
        book.close()
