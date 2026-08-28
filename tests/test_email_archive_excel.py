from __future__ import annotations

import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook

from main_service.codex_runner import CodexResult
from main_service.email_archive_excel import (
    EMPTY_VALUE,
    EXCEL_FIELDS,
    METADATA_TITLE,
    WORKSHEET_TITLE,
    EmailArchiveExcelStore,
    archive_save_candidates,
    natural_language_items,
)


def archive_result(**changes):
    result = {
        "발신자": "담당자",
        "날짜": "2026-08-22",
        "Topic": "지연이자 논의",
        "금액": 1234567,
        "통화": "KRW",
        "Business Impact": {
            "confirmed": ["공급사에 기준을 회신해야 합니다."],
            "estimated": ["금액이 달라질 수 있습니다."],
        },
        "Thread 진행 중 변경된 내용": [
            {"field": "회신 기한", "from": None, "to": "2026-08-17 18:00"}
        ],
        "결정된 내용": ["영업일 기준으로 산정하기로 했습니다."],
        "Open Item": ["최종 금액을 확인해야 합니다."],
    }
    result.update(changes)
    return result


class NaturalLanguageTests(unittest.TestCase):
    def test_flattens_dicts_and_lists_without_internal_keys(self):
        items = natural_language_items(archive_result()["Business Impact"])
        self.assertEqual(
            items,
            [
                "공급사에 기준을 회신해야 합니다.",
                "확인 필요: 금액이 달라질 수 있습니다.",
            ],
        )
        rendered = "\n".join(items)
        for hidden in ("confirmed", "estimated", "{", "}", "[", "]"):
            self.assertNotIn(hidden, rendered)

    def test_empty_values_have_one_friendly_fallback(self):
        for value in (None, "", [], {}):
            self.assertEqual(natural_language_items(value), [])
        self.assertEqual(EMPTY_VALUE, "확인된 내용 없음")

    def test_removes_duplicate_sentences(self):
        value = {"confirmed": ["같은 문장"], "estimated": ["같은 문장"]}
        self.assertEqual(natural_language_items(value), ["같은 문장"])

    def test_formats_change_objects_as_sentences(self):
        items = natural_language_items(archive_result()["Thread 진행 중 변경된 내용"])
        self.assertEqual(items, ["회신 기한이(가) 2026-08-17 18:00(으)로 확정되었습니다."])
        self.assertNotIn("from", items[0])
        self.assertNotIn("to", items[0])


class ExcelStoreTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path = Path(self.temp.name) / "exports" / "email_archive_results.xlsx"
        self.store = EmailArchiveExcelStore(self.path)
        self.email = {
            "thread_id": "thread-1",
            "received_at": "2026-08-22T10:00:00+09:00",
        }

    def tearDown(self):
        self.temp.cleanup()

    def test_appends_without_overwriting_existing_ab_rows_or_sheets(self):
        self.path.parent.mkdir(parents=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = WORKSHEET_TITLE
        sheet.append(EXCEL_FIELDS)
        sheet.append(("기존 발신자", date(2026, 8, 1), "기존 Topic", 10, "KRW", "기존", "기존", "기존", "기존"))
        workbook.create_sheet("기존 worksheet")["A1"] = "유지"
        workbook.save(self.path)

        saved = self.store.append(self.email, archive_result())

        self.assertTrue(saved.added)
        self.assertEqual(saved.row_count, 2)
        workbook = load_workbook(self.path)
        self.assertIn("기존 worksheet", workbook.sheetnames)
        self.assertEqual(workbook["기존 worksheet"]["A1"].value, "유지")
        sheet = workbook[WORKSHEET_TITLE]
        self.assertEqual(sheet.cell(row=2, column=1).value, "기존 발신자")
        self.assertEqual(sheet.cell(row=3, column=4).value, 1234567)
        self.assertIn("\n", sheet.cell(row=3, column=6).value)
        self.assertEqual(workbook[METADATA_TITLE].sheet_state, "hidden")
        workbook.close()

        # Main Service는 실험 모듈을 import하지 않지만, 기존 A/B reader도 결과를 열 수 있다.
        from experiments.ablation.skill_ab_test.excel_export import WithSkillExcelStore

        ab_store = WithSkillExcelStore(self.path, self.path.parent)
        self.assertIsNotNone(ab_store.read_bytes())

    def test_duplicate_is_blocked_but_changed_result_is_a_new_version(self):
        first = self.store.append(self.email, archive_result())
        duplicate = self.store.append(self.email, archive_result())
        changed = self.store.append(
            self.email, archive_result(**{"Open Item": ["새 확인 사항"]})
        )
        self.assertTrue(first.added)
        self.assertFalse(duplicate.added)
        self.assertTrue(changed.added)
        self.assertEqual(changed.row_count, 2)

    def test_new_latest_message_date_is_a_new_version(self):
        self.store.append(self.email, archive_result())
        newer_email = {**self.email, "received_at": "2026-08-23T09:00:00+09:00"}
        saved = self.store.append(newer_email, archive_result())
        self.assertTrue(saved.added)
        self.assertEqual(saved.row_count, 2)

    def test_one_save_appends_multiple_new_results_and_skips_duplicates(self):
        self.store.append(self.email, archive_result())
        second_email = {"thread_id": "thread-2", "received_at": "2026-08-23"}
        saved = self.store.append_many(
            [
                (self.email, archive_result()),
                (second_email, archive_result(Topic="새 논의")),
            ]
        )
        self.assertEqual(saved.added_count, 1)
        self.assertEqual(saved.duplicate_count, 1)
        self.assertEqual(saved.row_count, 2)

    def test_download_read_does_not_modify_workbook(self):
        self.store.append(self.email, archive_result())
        before = self.path.read_bytes()
        first = self.store.read_bytes()
        second = self.store.read_bytes()
        self.assertEqual(first, before)
        self.assertEqual(second, before)
        self.assertEqual(self.path.read_bytes(), before)

    def test_duplicate_only_save_does_not_modify_workbook(self):
        self.store.append(self.email, archive_result())
        before = self.path.read_bytes()
        saved = self.store.append_many([(self.email, archive_result())])
        self.assertEqual(saved.added_count, 0)
        self.assertEqual(saved.duplicate_count, 1)
        self.assertEqual(self.path.read_bytes(), before)

    def test_noncanonical_result_does_not_create_or_modify_workbook(self):
        with self.assertRaises(ValueError):
            self.store.append_many(
                [(self.email, {"archive_result": archive_result()})]
            )
        self.assertFalse(self.path.exists())


class SaveScopeTests(unittest.TestCase):
    def test_only_current_successful_analyzed_results_are_candidates(self):
        records = {
            "success": {
                "status": "ok",
                "email": {"thread_id": "selected"},
                "result": CodexResult(text="{}", parsed=archive_result()),
            },
            "failed": {
                "status": "error",
                "email": {"thread_id": "failed"},
                "error": "failed",
            },
            "unselected": {
                "status": "ok",
                "email": {"thread_id": "unselected"},
                "result": CodexResult(text="{}", parsed=archive_result()),
            },
        }
        candidates, failed = archive_save_candidates(
            records, ["success", "failed", "not-analyzed"]
        )
        self.assertEqual(len(candidates), 1)
        self.assertEqual(candidates[0][0]["thread_id"], "selected")
        self.assertEqual(failed, 1)

    def test_nested_or_noncanonical_result_is_not_an_excel_candidate(self):
        records = {
            "nested": {
                "status": "ok",
                "email": {"thread_id": "selected"},
                "result": CodexResult(
                    text="{}", parsed={"archive_result": archive_result()}
                ),
            }
        }
        candidates, failed = archive_save_candidates(records, ["nested"])
        self.assertEqual(candidates, [])
        self.assertEqual(failed, 1)


class ArchiveUiSourceTests(unittest.TestCase):
    def test_save_and_download_buttons_share_columns_and_have_exact_labels(self):
        source = (
            Path(__file__).resolve().parents[1] / "main_service" / "streamlit_app.py"
        ).read_text(encoding="utf-8")
        self.assertIn("save_column, download_column = st.columns(2)", source)
        self.assertIn('if save_column.button(\n        "💾 Save"', source)
        self.assertIn('download_column.download_button(\n        "📥 Download"', source)
        self.assertNotIn("st.json(value", source)
        self.assertIn(
            "Save: 분석이 완료된 미저장 아카이빙 결과를 Master Excel에 추가",
            source,
        )
        self.assertIn("새로 저장할 아카이빙 결과가 없습니다.", source)

    def test_multi_select_controls_and_dynamic_labels_are_present(self):
        source = (
            Path(__file__).resolve().parents[1] / "main_service" / "streamlit_app.py"
        ).read_text(encoding="utf-8")
        self.assertNotIn("archive_check.checkbox(", source)
        self.assertNotIn("st.checkbox(\n                    \"논의 아카이빙 선택\"", source)
        self.assertNotIn('"논의 아카이빙 선택"', source)
        self.assertNotIn('"전체 선택"', source)
        self.assertNotIn('"전체 해제"', source)
        self.assertIn('key=f"archive-pick::{inbox_key}::{case_id}"', source)
        self.assertIn('type="primary" if is_archive_picked else "secondary"', source)
        self.assertIn('if is_archive_picked\n                        else ":material/circle:"', source)
        self.assertIn('return "🗂️ 아카이빙할 메일 선택"', source)
        self.assertIn('return f"🗂️ 선택 {selected_count}건 아카이빙"', source)
        self.assertIn("disabled=not ready or not selected_items or running", source)

    def test_archive_button_is_in_unsupported_stage_header_not_page_top(self):
        source = (
            Path(__file__).resolve().parents[1] / "main_service" / "streamlit_app.py"
        ).read_text(encoding="utf-8")
        self.assertNotIn("render_archive_batch_controls", source)
        self.assertIn('if stage == "unsupported":\n                render_archive_stage_header', source)
        self.assertIn("title_column, run_column = st.columns", source)
        self.assertIn("title_column.subheader", source)
        self.assertIn("if run_column.button(", source)
        self.assertLess(source.index("title_column.subheader"), source.index("if run_column.button("))

    def test_results_and_excel_controls_are_below_unsupported_cards(self):
        source = (
            Path(__file__).resolve().parents[1] / "main_service" / "streamlit_app.py"
        ).read_text(encoding="utf-8")
        cards = source.index("render_fold_list(bucket, settings)")
        results = source.index("render_archive_batch_results()", cards)
        self.assertLess(cards, results)


if __name__ == "__main__":
    unittest.main()
