from __future__ import annotations

import json
import os
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from .models import (
    ActionItem,
    ArchiveRecord,
    Attachment,
    FieldChange,
    ImportantNumber,
    SourceEmail,
)

RECORDS = "Records"
ACTION_ITEMS = "Action Items"
ATTACHMENTS = "Attachments"
SOURCE_EMAILS = "Source Emails"
CHANGE_HISTORY = "Change History"

HEADERS = {
    RECORDS: [
        "Record ID",
        "Title",
        "Requester",
        "Participants",
        "Start Date",
        "Latest Update Date",
        "Thread Summary",
        "Key Points",
        "Decisions",
        "Important Numbers",
        "Open Items",
        "Current Status",
        "Source Thread ID",
        "Reference IDs",
        "Last Updated",
    ],
    ACTION_ITEMS: [
        "Record ID",
        "Action ID",
        "Description",
        "Owner",
        "Due Date",
        "Status",
        "Source Message ID",
    ],
    ATTACHMENTS: [
        "Record ID",
        "Attachment ID",
        "File Name",
        "File Type",
        "MIME Type",
        "Access Status",
        "Summary",
        "Source Message ID",
    ],
    SOURCE_EMAILS: ["Record ID", "Message ID", "Thread ID", "Subject", "Sent At", "Source Type"],
    CHANGE_HISTORY: [
        "Change ID",
        "Record ID",
        "Field",
        "Change Type",
        "Previous Value",
        "New Value",
        "Changed At",
        "Source Message ID",
        "Reason / Evidence",
    ],
}


class ExcelRecordRepository:
    """Excel-backed record repository with atomic, validated replacement."""

    def __init__(self, workbook_path: str | Path, repository_root: str | Path) -> None:
        self.repository_root = Path(repository_root).resolve()
        self.workbook_path = Path(workbook_path).resolve()
        if self.workbook_path.suffix.lower() != ".xlsx":
            raise ValueError("Archive storage must be an .xlsx file.")
        try:
            self.workbook_path.relative_to(self.repository_root)
        except ValueError as error:
            raise ValueError("Archive workbook must be inside the repository.") from error

    def list_records(self) -> list[ArchiveRecord]:
        if not self.workbook_path.exists():
            return []
        workbook = load_workbook(self.workbook_path, data_only=True)
        try:
            self._validate_structure(workbook)
            actions = self._group_rows(workbook[ACTION_ITEMS])
            attachments = self._group_rows(workbook[ATTACHMENTS])
            sources = self._group_rows(workbook[SOURCE_EMAILS])
            records: list[ArchiveRecord] = []
            for row in self._dict_rows(workbook[RECORDS]):
                record_id = _string(row["Record ID"])
                records.append(
                    ArchiveRecord(
                        record_id=record_id,
                        record_id_candidate=record_id,
                        title=_string(row["Title"]),
                        requester=_string(row["Requester"]),
                        participants=_json_list(row["Participants"]),
                        start_date=_string(row["Start Date"]),
                        latest_update_date=_string(row["Latest Update Date"]),
                        thread_summary=_string(row["Thread Summary"]),
                        key_points=_json_list(row["Key Points"]),
                        decisions=_json_list(row["Decisions"]),
                        important_numbers=[ImportantNumber.model_validate(item) for item in _json_list(row["Important Numbers"])],
                        open_items=_json_list(row["Open Items"]),
                        current_status=_string(row["Current Status"]),
                        reference_ids=_json_list(row["Reference IDs"]),
                        last_updated=_string(row["Last Updated"]),
                        action_items=[self._action_from_row(item) for item in actions.get(record_id, [])],
                        attachments=[self._attachment_from_row(item) for item in attachments.get(record_id, [])],
                        source_emails=[self._source_from_row(item) for item in sources.get(record_id, [])],
                    )
                )
            for record in records:
                record.due_dates = _unique(
                    item.due_date for item in record.action_items if item.due_date != "unknown"
                )
            return records
        finally:
            workbook.close()

    def get_record(self, record_id: str) -> ArchiveRecord | None:
        return next((record for record in self.list_records() if record.record_id == record_id), None)

    def list_change_history(self, record_id: str) -> list[FieldChange]:
        if not self.workbook_path.exists():
            return []
        workbook = load_workbook(self.workbook_path, data_only=True, read_only=True)
        try:
            self._validate_structure(workbook)
            history: list[FieldChange] = []
            for row in self._dict_rows(workbook[CHANGE_HISTORY]):
                if _string(row["Record ID"]) != record_id:
                    continue
                history.append(
                    FieldChange(
                        field_path=_string(row["Field"]),
                        change_type=_string(row["Change Type"]),
                        previous_value=_parse_json_value(row["Previous Value"]),
                        new_value=_parse_json_value(row["New Value"]),
                        changed_at=_string(row["Changed At"]),
                        source_message_id=_string(row["Source Message ID"]),
                        reason=_string(row["Reason / Evidence"]),
                    )
                )
            return history
        finally:
            workbook.close()

    def save(self, record: ArchiveRecord, changes: list[FieldChange], operation: str) -> None:
        self.workbook_path.parent.mkdir(parents=True, exist_ok=True)
        workbook = self._load_or_create()
        temp_path: Path | None = None
        try:
            self._upsert_record(workbook[RECORDS], record)
            self._replace_child_rows(workbook[ACTION_ITEMS], record.record_id, self._action_rows(record))
            self._replace_child_rows(workbook[ATTACHMENTS], record.record_id, self._attachment_rows(record))
            self._replace_child_rows(workbook[SOURCE_EMAILS], record.record_id, self._source_rows(record))
            self._append_changes(workbook[CHANGE_HISTORY], record.record_id, changes)

            with tempfile.NamedTemporaryFile(
                prefix=f".{self.workbook_path.stem}.",
                suffix=".xlsx",
                dir=self.workbook_path.parent,
                delete=False,
            ) as handle:
                temp_path = Path(handle.name)
            workbook.save(temp_path)
            self._validate_saved_file(temp_path, record.record_id)
            os.replace(temp_path, self.workbook_path)
            temp_path = None
        finally:
            workbook.close()
            if temp_path is not None:
                temp_path.unlink(missing_ok=True)

    def _load_or_create(self):
        if self.workbook_path.exists():
            workbook = load_workbook(self.workbook_path)
            self._validate_structure(workbook)
            return workbook
        workbook = Workbook()
        default = workbook.active
        workbook.remove(default)
        for sheet_name, headers in HEADERS.items():
            sheet = workbook.create_sheet(sheet_name)
            sheet.append(headers)
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = f"A1:{_column_letter(len(headers))}1"
        return workbook

    def _validate_structure(self, workbook) -> None:
        missing = [name for name in HEADERS if name not in workbook.sheetnames]
        if missing:
            raise ValueError(f"Archive workbook is missing sheets: {', '.join(missing)}")
        for name, expected in HEADERS.items():
            actual = [cell.value for cell in workbook[name][1]]
            if actual != expected:
                raise ValueError(f"Archive workbook sheet '{name}' has an incompatible header.")

    def _validate_saved_file(self, path: Path, record_id: str) -> None:
        candidate = load_workbook(path, read_only=True, data_only=True)
        try:
            self._validate_structure(candidate)
            ids = {_string(row[0]) for row in candidate[RECORDS].iter_rows(min_row=2, values_only=True)}
            if record_id not in ids:
                raise ValueError("Saved workbook validation could not find the record ID.")
        finally:
            candidate.close()

    def _upsert_record(self, sheet: Worksheet, record: ArchiveRecord) -> None:
        values = [
            record.record_id,
            record.title,
            record.requester,
            _json(record.participants),
            record.start_date,
            record.latest_update_date,
            record.thread_summary,
            _json(record.key_points),
            _json(record.decisions),
            _json([item.model_dump(mode="json") for item in record.important_numbers]),
            _json(record.open_items),
            record.current_status,
            record.source_emails[-1].thread_id if record.source_emails else "unknown",
            _json(record.reference_ids),
            record.last_updated,
        ]
        target_row = next(
            (index for index in range(2, sheet.max_row + 1) if _string(sheet.cell(index, 1).value) == record.record_id),
            None,
        )
        if target_row is None:
            sheet.append(values)
            return
        for column, value in enumerate(values, start=1):
            sheet.cell(target_row, column, value)

    @staticmethod
    def _replace_child_rows(sheet: Worksheet, record_id: str, rows: list[list[Any]]) -> None:
        for row_number in range(sheet.max_row, 1, -1):
            if _string(sheet.cell(row_number, 1).value) == record_id:
                sheet.delete_rows(row_number)
        for row in rows:
            sheet.append(row)

    @staticmethod
    def _append_changes(sheet: Worksheet, record_id: str, changes: list[FieldChange]) -> None:
        existing_ids = {_string(sheet.cell(row, 1).value) for row in range(2, sheet.max_row + 1)}
        for index, change in enumerate(changes, start=1):
            base = f"{record_id}|{change.field_path}|{change.changed_at}|{change.source_message_id}|{index}"
            change_id = _stable_id("CHG", base)
            if change_id in existing_ids:
                continue
            sheet.append(
                [
                    change_id,
                    record_id,
                    change.field_path,
                    change.change_type,
                    _json_value(change.previous_value),
                    _json_value(change.new_value),
                    change.changed_at,
                    change.source_message_id,
                    change.reason,
                ]
            )

    @staticmethod
    def _action_rows(record: ArchiveRecord) -> list[list[Any]]:
        return [
            [
                record.record_id,
                item.action_id_candidate,
                item.description,
                item.owner,
                item.due_date,
                item.status,
                item.source_message_id,
            ]
            for item in record.action_items
        ]

    @staticmethod
    def _attachment_rows(record: ArchiveRecord) -> list[list[Any]]:
        return [
            [
                record.record_id,
                item.attachment_id,
                item.file_name,
                item.file_type,
                item.mime_type,
                item.access_status,
                item.summary,
                item.source_message_id,
            ]
            for item in record.attachments
        ]

    @staticmethod
    def _source_rows(record: ArchiveRecord) -> list[list[Any]]:
        return [
            [record.record_id, item.message_id, item.thread_id, item.subject, item.sent_at, item.source_type]
            for item in record.source_emails
        ]

    @staticmethod
    def _dict_rows(sheet: Worksheet) -> list[dict[str, Any]]:
        headers = [cell.value for cell in sheet[1]]
        return [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True) if row[0]]

    def _group_rows(self, sheet: Worksheet) -> dict[str, list[dict[str, Any]]]:
        grouped: dict[str, list[dict[str, Any]]] = {}
        for row in self._dict_rows(sheet):
            grouped.setdefault(_string(row["Record ID"]), []).append(row)
        return grouped

    @staticmethod
    def _action_from_row(row: dict[str, Any]) -> ActionItem:
        return ActionItem(
            action_id_candidate=_string(row["Action ID"]),
            description=_string(row["Description"]),
            owner=_string(row["Owner"]),
            due_date=_string(row["Due Date"]),
            status=_string(row["Status"]),
            source_message_id=_string(row["Source Message ID"]),
        )

    @staticmethod
    def _attachment_from_row(row: dict[str, Any]) -> Attachment:
        return Attachment(
            attachment_id=_string(row["Attachment ID"]),
            file_name=_string(row["File Name"]),
            file_type=_string(row["File Type"]),
            mime_type=_string(row["MIME Type"]),
            access_status=_string(row["Access Status"]),
            summary=_string(row["Summary"]),
            source_message_id=_string(row["Source Message ID"]),
        )

    @staticmethod
    def _source_from_row(row: dict[str, Any]) -> SourceEmail:
        return SourceEmail(
            message_id=_string(row["Message ID"]),
            thread_id=_string(row["Thread ID"]),
            subject=_string(row["Subject"]),
            sent_at=_string(row["Sent At"]),
            source_type=_string(row["Source Type"]),
        )


def _string(value: Any) -> str:
    return "unknown" if value is None or value == "" else str(value)


def _json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True)


def _json_value(value: Any) -> str:
    return value if isinstance(value, str) else _json(value)


def _json_list(value: Any) -> list[Any]:
    if value is None or value == "":
        return []
    parsed = json.loads(str(value))
    if not isinstance(parsed, list):
        raise ValueError("Expected a JSON list in archive workbook.")
    return parsed


def _parse_json_value(value: Any) -> Any:
    if value is None:
        return "unknown"
    if not isinstance(value, str):
        return value
    try:
        return json.loads(value)
    except json.JSONDecodeError:
        return value


def _unique(values) -> list[str]:
    return list(dict.fromkeys(values))


def _column_letter(number: int) -> str:
    result = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _stable_id(prefix: str, value: str) -> str:
    import hashlib

    return f"{prefix}-{hashlib.sha256(value.encode('utf-8')).hexdigest()[:16].upper()}"


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()
